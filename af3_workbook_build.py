#!/usr/bin/env python3
"""
Every raw AlphaFold Server value and every calculated table, in one Excel workbook.

    python3 af3_workbook_build.py [results_dir]        # default: data/ next to this script

Writes af3_idr_workbook.xlsx. The main sheet, Jobs, is af3_idr_pairs_1.csv with its
columns unchanged and the raw and calculated values appended; Models has one row per
job x server sample under the same leading columns. Header colour gives each column's
source: grey = pair table, blue = raw server output, green = calculated.

Per-sample metrics are recomputed with parse_af3_results' own functions, and the top
sample and sample spread must reproduce af3_scored.csv before anything is written, so
the workbook cannot drift from the pipeline. PAE and contact-probability matrices are
too big for a sheet; view them with viewer/serve.py.

Requires numpy and openpyxl.
"""
import csv, datetime, glob, json, math, os, re, statistics, sys
from multiprocessing import Pool

sys.dont_write_bytecode = True                    # never litter the repo with .pyc
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import parse_af3_results as P
import baseline_leakage as B

OUT = os.path.join(HERE, 'af3_idr_workbook.xlsx')
MEM = {'P06', 'P07', 'P20', 'P21'}                # memorised pipeline controls
ID_COLS = ['job_id', 'job_name', 'chain_A', 'chain_B', 'label', 'pair_type',
           'interaction_class', 'binding_mode', 'llps', 'in_pdb', 'neg_type',
           'evidence_strength', 'stage', 'len_A', 'len_B', 'total_residues']
PANEL = [m for m, _ in P.METRICS]
LOCAL = [m for m, _ in P.EXPLORATORY]
SIGN = dict(P.METRICS + P.EXPLORATORY)
RAW_SCALARS = ['ranking_score', 'iptm', 'ptm', 'fraction_disordered', 'has_clash']
# everything the parser derives for one sample, minus the scalars it only copies
SAMPLE_CALC = [m for m in PANEL + LOCAL if m not in RAW_SCALARS] + [
    'pae_inter_min', 'chain_ptm_mean', 'plddt_chainA_mean', 'plddt_chainB_mean']
FILL = {'input': 'D9D9D9', 'raw': 'DDEBF7', 'calc': 'E2EFDA'}

DEFINITIONS = {
    'iptm': 'interface pTM; the pre-registered primary metric',
    'chain_pair_iptm': 'mean of the off-diagonal chain_pair_iptm entries',
    'ranking_score': "AF3's ranking score; it picks the top sample",
    'ptm': 'pTM of the whole complex',
    'fraction_disordered': 'fraction of the structure AF3 calls disordered',
    'chain_pair_pae_min': 'smallest off-diagonal chain_pair_pae_min (Å)',
    'pae_inter_mean': 'mean PAE over every interchain token pair, both directions (Å)',
    'pae_inter_p05': '5th percentile of interchain PAE (Å)',
    'pae_inter_frac_lt10': f'fraction of interchain PAE below {P.PAE_CONF:g} Å',
    'pae_inter_frac_lt5': f'fraction of interchain PAE below {P.PAE_TIGHT:g} Å',
    'contact_prob_sum': 'sum of the A×B contact-probability block',
    'contact_prob_density': 'contact_prob_sum / (len_A × len_B)',
    'contact_prob_max': 'largest A×B contact probability',
    'n_contacts_p50': f'A×B token pairs with contact probability > {P.CONTACT_P:g}',
    'contact_participation_ratio': '(Σp)² / Σp² over the A×B block: ~1 = one contact, large = many weak ones',
    'plddt_mean': 'mean atom pLDDT',
    'ipsae': f'ipSAE (Dunbrack 2025), PAE cutoff {P.IPSAE_CUTOFF:g} Å; exploratory, added after stage 1',
    'lis': f'Local Interaction Score (Kim et al. 2024), PAE cutoff {P.LIS_CUTOFF:g} Å; exploratory, added after stage 1',
}

# ---------------------------------------------------------------- inputs

def read_csv(name):
    with open(os.path.join(HERE, name), newline='') as f:
        return list(csv.DictReader(f))


def load_json(path):
    with open(path) as f:
        return json.load(f)


# ---------------------------------------------------------------- result folders
# viewer/serve.py uses these too

def find_job_dirs(results_dir):
    """Normalised folder name -> path, for every folder holding server output."""
    dirs = {}
    for root, _dirs, files in os.walk(results_dir):
        if any('_summary_confidences_' in f for f in files):
            dirs[P.norm(os.path.basename(root))] = root
    return dirs


def match_dir(job_name, dirs):
    """The parser's rule: exact normalised job name, else a unique prefix."""
    key = P.norm(job_name)
    if key in dirs:
        return dirs[key]
    hits = [k for k in dirs if k.startswith(key)]
    return dirs[hits[0]] if len(hits) == 1 else None


def model_files(job_dir, kind):
    """{sample index: path} for kind 'summary_confidences' or 'full_data'."""
    out = {}
    for f in glob.glob(os.path.join(job_dir, f'*_{kind}_*.json')):
        m = re.search(r'_(\d+)\.json$', f)
        if m:
            out[int(m.group(1))] = f
    return dict(sorted(out.items()))


def top_model(ranking_scores):
    """Position of AF3's best sample, picked exactly as parse_af3_results.main picks it."""
    return max(range(len(ranking_scores)), key=lambda i: ranking_scores[i] or -1e9)


# ---------------------------------------------------------------- per job

def read_samples(job_dir):
    """[(raw, calc)] per server sample: summary_confidences as written, parser metrics."""
    full = model_files(job_dir, 'full_data')
    out = []
    for i, path in model_files(job_dir, 'summary_confidences').items():
        s = load_json(path)
        chains = list(dict.fromkeys(s['chain_ids']))
        raw = {'model': i, **{k: s.get(k) for k in RAW_SCALARS}}
        for vec in ('chain_iptm', 'chain_ptm'):
            raw.update({f'{vec}_{c}': v for c, v in zip(chains, s.get(vec) or [])})
        for mat in ('chain_pair_iptm', 'chain_pair_pae_min'):
            for c1, row in zip(chains, s.get(mat) or []):
                raw.update({f'{mat}_{c1}{c2}': v for c2, v in zip(chains, row)})
        calc = P.summary_metrics(s)
        if i in full:
            fd = load_json(full[i])
            raw['n_tokens'], raw['n_atoms'] = len(fd['token_chain_ids']), len(fd['atom_plddts'])
            calc.update(P.interchain_metrics(fd))
        out.append((raw, calc))
    return out


def job_samples(item):
    job_id, job_dir = item
    return job_id, read_samples(job_dir)


def as_scored(calcs):
    """(top position, the af3_scored.csv cells parse_af3_results.main writes for these samples)"""
    best = top_model([c.get('ranking_score') for c in calcs])
    cells = {}
    for name in PANEL + LOCAL:
        vals = [c[name] for c in calcs if c.get(name) is not None]
        if not vals:
            cells[name] = cells[f'{name}_sd'] = ''
            continue
        cells[name] = str(round(calcs[best].get(name, vals[0]), 5))
        cells[f'{name}_sd'] = str(round(statistics.pstdev(vals), 5)) if len(vals) > 1 else ''
    return best, cells


def request_info(job_dir, pair):
    """What the server ran, from its job_request.json (the batch files are not the record)."""
    found = glob.glob(os.path.join(job_dir, '*_job_request.json'))
    if not found:
        return {}
    req = load_json(found[0])
    req = req[0] if isinstance(req, list) else req
    chains = [e['proteinChain'] for e in req.get('sequences', []) if 'proteinChain' in e]
    ran = [c['sequence'] for c in chains for _ in range(int(c.get('count', 1)))]
    templates = {c.get('useStructureTemplate') for c in chains}
    return {'model_seeds': ', '.join(req.get('modelSeeds') or []),
            'chain_counts': json.dumps([int(c.get('count', 1)) for c in chains]),   # [2] = homodimer entry
            'use_structure_template': templates.pop() if len(templates) == 1 else 'mixed',
            'seqs_match_pair_table': ran == [pair['seq_A'], pair['seq_B']]}


def msa_info(entry):
    """MSA depth and template hits per chain; a homodimer's chains share one entry ('b_a')."""
    def of(chain, items):
        return [v for k, v in items if chain in k.split('/')[0].split('_')]
    out = {}
    for chain in ('A', 'B'):
        c = chain.lower()
        for kind in ('paired', 'unpaired'):
            depth = of(c, [(k, v) for k, v in entry.get('msa', {}).items() if k.endswith('/' + kind)])
            out[f'msa_{kind}_{chain}'] = depth[0] if depth else None
        hits = of(c, entry.get('templates', {}).items())
        out[f'templates_{chain}'] = ', '.join(hits[0]) if hits else None
    return out


# ---------------------------------------------------------------- sheets
# each returns (columns, rows, column -> 'input' | 'raw' | 'calc')

def calc_only(_col):
    return 'calc'


def jobs_sheet(pairs, scored, report, samples, tops, located):
    table = {r['job_id']: r for r in report['table']}
    scored_cols = list(next(iter(scored.values())))
    metric_cols = scored_cols[scored_cols.index('n_seeds'):]
    extra_cols = ['baseline_pooled', 'baseline_arm', 'has_clash_any', 'chain_ptm_mean']
    input_cols = list(pairs[0])
    raw_cols = ['status', 'result_folder', 'model_seeds', 'chain_counts',
                'use_structure_template', 'seqs_match_pair_table', 'msa_paired_A',
                'msa_unpaired_A', 'msa_paired_B', 'msa_unpaired_B', 'templates_A', 'templates_B']
    rows = []
    for p in pairs:
        j, d = p['job_id'], located[p['job_id']]
        row = dict(p, status='scored' if j in scored else 'no results', memorised_control=j in MEM)
        if d:
            row['result_folder'] = os.path.relpath(d, HERE)
            row.update(request_info(d, p))
            row.update(msa_info(report['msa_templates'].get(j, {})))
        if j in samples:
            row['top_model'] = samples[j][tops[j]][0]['model']
        row.update({k: scored.get(j, {}).get(k) for k in metric_cols})
        row.update({k: table.get(j, {}).get(k) for k in extra_cols})
        rows.append(row)
    cols = input_cols + raw_cols + ['memorised_control', 'top_model'] + metric_cols + extra_cols
    return cols, rows, lambda c: 'input' if c in input_cols else 'raw' if c in raw_cols else 'calc'


def models_sheet(pairs, samples, tops):
    rows, raw_cols = [], []
    for p in pairs:
        s = samples.get(p['job_id'])
        if not s:
            continue
        scores = [calc.get('ranking_score') for _raw, calc in s]
        order = sorted(range(len(s)), key=lambda i: (-(scores[i] or -1e9), i))
        for i, (raw, calc) in enumerate(s):
            raw_cols += [k for k in raw if k not in raw_cols]
            rows.append({**{k: p[k] for k in ID_COLS}, **raw,
                         'rank': order.index(i) + 1, 'is_top_model': i == tops[p['job_id']],
                         **{k: calc.get(k) for k in SAMPLE_CALC}})
    cols = ID_COLS + raw_cols + ['rank', 'is_top_model'] + SAMPLE_CALC
    return cols, rows, lambda c: 'input' if c in ID_COLS else 'raw' if c in raw_cols else 'calc'


def composition_sheets(pairs, scored, report):
    """Per-job composition features and single-feature AUCs, checked against the report."""
    table = {r['job_id']: r for r in report['table']}
    feats = {p['job_id']: B.pair_feats(p['seq_A'], p['seq_B']) for p in pairs}
    names = sorted(next(iter(feats.values())))
    rows = []
    for p in pairs:
        t = table.get(p['job_id'], {})
        rows.append({**{k: p[k] for k in ID_COLS}, **feats[p['job_id']],
                     'baseline_pooled': t.get('baseline_pooled'), 'baseline_arm': t.get('baseline_arm')})
    per_job = (ID_COLS + names + ['baseline_pooled', 'baseline_arm'], rows,
               lambda c: 'input' if c in ID_COLS else 'calc')

    # the subsets af3_report_stats.py fits the baseline on: scored jobs with real sequences
    real = [p for p in pairs if p['job_id'] in scored and p['label'] != 'negative_scramble']
    arms = [('all', 'pooled', real),
            ('homotypic', 'homotypic', [p for p in real if p['pair_type'] == 'homotypic']),
            ('heterotypic', 'heterotypic', [p for p in real if p['pair_type'] == 'heterotypic'])]
    loo, rows, problems = report['meta']['baseline_loo'], [], []
    for arm, key, sub in arms:
        pos = [feats[p['job_id']] for p in sub if p['label'] == 'positive']
        neg = [feats[p['job_id']] for p in sub if p['label'] != 'positive']
        n = {'arm': arm, 'n_pos': len(pos), 'n_neg': len(neg)}
        rows.append({**n, 'feature': 'LOO logistic, all features',
                     'auc_raw': loo[f'{key}_raw'], 'auc_two_sided': loo[f'{key}_two_sided']})
        single = []
        for name in names:
            a = P.auc([f[name] for f in pos], [f[name] for f in neg])
            single.append({**n, 'feature': name, 'auc_raw': a, 'auc_two_sided': B.two_sided(a)})
        single.sort(key=lambda d: -d['auc_two_sided'])
        if any(d['feature'] != w['feature'] or abs(d['auc_two_sided'] - w['auc_two_sided']) > 1e-12
               for d, w in zip(single, report['single_features'][arm])):
            problems.append(f'composition feature AUCs ({arm}) differ from af3_report_data.json')
        rows.extend(single)
    aucs = (['arm', 'feature', 'n_pos', 'n_neg', 'auc_raw', 'auc_two_sided'], rows, calc_only)
    return per_job, aucs, problems


def groups_sheet(report):
    rows = []
    for r in report['refs']:
        row = {k: r.get(k) for k in ('family', 'group', 'note', 'n_pos', 'n_neg')}
        row.update(baseline=r.get('base_key'), baseline_auc=r.get('baseline_auc'),
                   length_only_auc=r.get('length_auc'))
        for side in ('pos', 'neg'):
            lo, hi = r.get(f'len_{side}_range') or (None, None)
            row.update({f'len_{side}_median': r.get(f'len_{side}_median'),
                        f'len_{side}_min': lo, f'len_{side}_max': hi})
        rows.append(row)
    return list(rows[0]), rows, calc_only


def metrics_sheet(report):
    local = report['local_pae']
    conf = {c['metric']: c for c in report['confounds'] + local['confounds']}
    scr = {**report['scramble_summary'], **local['scramble_summary']}
    matched = {}
    for m in report['matched_internal'] + local['matched_internal']:
        matched.setdefault(f"{m['positive']}_above_{'_'.join(m['negatives'])}", {}).update(m['correct'])
    rows = []
    for name in PANEL + LOCAL:
        s = scr[name]
        rows.append({
            'metric': name,
            'family': 'pre-registered panel' if name in PANEL else 'exploratory (after stage 1)',
            'primary': name == P.PRIMARY,
            'more_interaction_when': 'higher' if SIGN[name] > 0 else 'lower',
            'definition': DEFINITIONS[name],
            'rho_vs_total_residues': conf[name]['rho_length'],
            'length_confounded': conf[name]['flagged'],
            'scramble_parent_wins': s['wins'], 'scramble_n': s['n'], 'scramble_ties': s.get('ties'),
            'scramble_sign_p': s['sign_p'], 'scramble_median_delta': s['median_delta'],
            **{col: correct.get(name) for col, correct in matched.items()}})
    return list(rows[0]), rows, calc_only


def scramble_sheet(report):
    local = {s['scramble']: s for s in report['local_pae']['scramble_pairs']}
    rows = []
    for s in report['scramble_pairs']:
        row = {k: s[k] for k in ('scramble', 'parent', 'scramble_name', 'parent_name', 'memorised_parent')}
        parent = {**s['parent_vals'], **local[s['scramble']]['parent_vals']}
        scram = {**s['scramble_vals'], **local[s['scramble']]['scramble_vals']}
        for m in PANEL + LOCAL:
            row.update({f'{m}_parent': parent[m], f'{m}_scramble': scram[m],
                        f'{m}_delta': parent[m] - scram[m]})
        rows.append(row)
    return list(rows[0]), rows, calc_only


def class_sheet(report):
    rows = []
    for d in report['distributions']:
        row = {'class': d['class'], 'n': d['n'], 'jobs': ', '.join(d['jobs'])}
        for m in PANEL + ['total_residues']:
            row.update({f'{m}_{stat}': d.get(m, {}).get(stat) for stat in ('median', 'min', 'max')})
        rows.append(row)
    return list(rows[0]), rows, calc_only


# ---------------------------------------------------------------- writing

def typed(v):
    """Cell value: numbers as numbers - only text Python itself writes for a number, so an id
    like PDB 1e50 stays text - True/False as booleans, blanks empty, lists joined."""
    if isinstance(v, str):
        if v in ('True', 'False'):
            return v == 'True'
        for kind in (int, float):
            try:
                if repr(kind(v)) == v:
                    v = kind(v)
                    break
            except ValueError:
                pass
        else:
            return v or None
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, (list, tuple)):
        return ', '.join(map(str, v))
    return v


def shown(v):
    """Rough displayed width of a cell, for column sizing."""
    if v is None:
        return 0
    return len(f'{v:.5f}'.rstrip('0')) if isinstance(v, float) else len(str(v))


def write_workbook(path, sheets, intro, notes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = {k: PatternFill('solid', fgColor=c) for k, c in FILL.items()}
    bold = Font(bold=True)
    wb = Workbook()
    readme = wb.active
    readme.title = 'README'

    for sh in sheets:
        ws = wb.create_sheet(sh['title'])
        ws.append(sh['cols'])
        for r in sh['rows']:
            ws.append([typed(r.get(c)) for c in sh['cols']])
        width = [len(c) + 3 for c in sh['cols']]
        for row in ws.iter_rows(min_row=2):
            for j, cell in enumerate(row):
                if isinstance(cell.value, float):
                    cell.number_format = '0.00###'
                width[j] = max(width[j], shown(cell.value) + 2)
        for j, col in enumerate(sh['cols'], start=1):
            head = ws.cell(row=1, column=j)
            head.font, head.fill = bold, fill[sh['kind'](col)]
            ws.column_dimensions[get_column_letter(j)].width = min(width[j - 1], 48)
        ws.freeze_panes = sh['freeze']
        ws.auto_filter.ref = ws.dimensions

    def line(*cells, font=None, fill_key=None):
        readme.append(list(cells))
        r = readme.max_row
        for j in range(1, len(cells) + 1):
            if font:
                readme.cell(row=r, column=j).font = font
        if fill_key:
            readme.cell(row=r, column=1).fill = fill[fill_key]

    line('AF3 IDR–IDR benchmark: raw and calculated data', font=Font(bold=True, size=14))
    line(intro)
    line()
    line('Header colour', font=bold)
    line('grey', 'input: the curated pair and region tables', fill_key='input')
    line('blue', 'raw: what AlphaFold Server ran and returned', fill_key='raw')
    line('green', 'calculated: parse_af3_results.py, af3_report_stats.py, baseline_leakage.py',
         fill_key='calc')
    line()
    line('Sheet', 'Rows', 'One row per', 'Contents', font=bold)
    for sh in sheets:
        line(sh['title'], len(sh['rows']), sh['per'], sh['what'])
    line()
    line('Notes', font=bold)
    for n in notes:
        line(n)
    for col, w in zip('ABCD', (16, 8, 24, 60)):
        readme.column_dimensions[col].width = w
    wb.save(path)


# ---------------------------------------------------------------- main

def main(results_dir):
    pairs = read_csv('af3_idr_pairs_1.csv')
    scored = {r['job_id']: r for r in read_csv('af3_scored.csv')}
    report = load_json(os.path.join(HERE, 'af3_report_data.json'))

    dirs = find_job_dirs(results_dir)
    located = {p['job_id']: match_dir(p['job_name'], dirs) for p in pairs}
    todo = [(j, d) for j, d in located.items() if d]
    print(f'reading every sample in {len(todo)} result folders ...')
    with Pool() as pool:
        samples = dict(pool.map(job_samples, todo))

    # refuse to write anything that disagrees with the pipeline's own outputs
    problems, tops = [], {}
    if set(samples) != set(scored):
        problems.append(f'jobs {sorted(set(samples) ^ set(scored))} differ between results and af3_scored.csv')
    for j, s in samples.items():
        tops[j], cells = as_scored([calc for _raw, calc in s])
        off = [k for k, v in cells.items() if j in scored and scored[j].get(k) != v]
        if off:
            problems.append(f'{j}: {", ".join(off)} differ from af3_scored.csv')
    composition, composition_auc, comp_problems = composition_sheets(pairs, scored, report)
    problems += comp_problems
    if problems:
        sys.exit('workbook not written; rerun parse_af3_results.py and af3_report_stats.py:\n  '
                 + '\n  '.join(problems))

    def sheet(title, spec, freeze, per, what):
        cols, rows, kind = spec
        return dict(title=title, cols=cols, rows=rows, kind=kind, freeze=freeze, per=per, what=what)

    regions = read_csv('af3_idr_regions.csv')
    stats = read_csv('af3_all_statistics.csv')
    sheets = [
        sheet('Jobs', jobs_sheet(pairs, scored, report, samples, tops, located), 'C2', 'job',
              "af3_idr_pairs_1.csv as is, then what the server ran and returned, then the job's "
              'scored metrics from af3_scored.csv (top sample value, SD over samples) and its '
              'composition-baseline scores'),
        sheet('Models', models_sheet(pairs, samples, tops), 'C2', 'job × server sample',
              'every value in summary_confidences_N.json, then the parser metrics recomputed for '
              'that sample from full_data_N.json; rank 1 = top sample'),
        sheet('Regions', (list(regions[0]), regions, lambda c: 'input'), 'B2', 'IDR region',
              'af3_idr_regions.csv as is'),
        sheet('Composition', composition, 'C2', 'job',
              'the sequence-composition features baseline_leakage.pair_feats computes, and the '
              "job's leave-one-out baseline scores"),
        sheet('Composition_AUC', composition_auc, 'C2', 'arm × feature',
              'how well each composition feature alone, and the LOO logistic model on all of '
              'them, separates positives from negatives'),
        sheet('Statistics', (list(stats[0]), stats, calc_only), 'D2', 'grouping × metric',
              'af3_all_statistics.csv as is: AUC, bootstrap CI, permutation p, BH q, ΔAUC vs '
              'composition with CI'),
        sheet('Groups', groups_sheet(report), 'C2', 'grouping',
              'group sizes, composition-baseline AUC, length-only AUC and chain-length spread'),
        sheet('Metrics', metrics_sheet(report), 'B2', 'metric',
              'definition and direction, length correlation, scramble sign test, matched-internal '
              'orderings'),
        sheet('Scramble_pairs', scramble_sheet(report), 'C2', 'scramble / parent pair',
              'each metric for the parent, its scramble, and parent − scramble'),
        sheet('Class_summary', class_sheet(report), 'B2', 'job class',
              'median, min and max of each panel metric and of total_residues'),
    ]

    loo = report['meta']['baseline_loo']
    lower = [m for m in PANEL + LOCAL if SIGN[m] < 0]
    pending = [p['job_id'] for p in pairs if p['job_id'] not in samples]
    biggest = max(raw.get('n_tokens', 0) for s in samples.values() for raw, _calc in s)
    today = datetime.date.today()
    intro = (f'Built {today.day} {today:%B %Y} by af3_workbook_build.py from af3_idr_pairs_1.csv, '
             f'af3_idr_regions.csv, the {len(samples)} result folders under '
             f'{os.path.relpath(results_dir, HERE)}/, af3_scored.csv, af3_report_data.json and '
             'af3_all_statistics.csv. Rebuild it rather than editing it.')
    notes = [
        f'{len(samples)} of {len(pairs)} jobs have results. {", ".join(pending) or "None"} '
        'carry pair-table columns only.',
        "Top sample: the highest ranking_score among a job's server samples, first on ties, as "
        "parse_af3_results.py picks it. Job-level metrics are that sample's values; *_sd is the "
        "population SD over the samples. n_seeds counts samples, which share one seed (model_seeds).",
        f'Checked at build time: every sample recomputed from full_data reproduces af3_scored.csv for '
        f'all {len(samples)} jobs, and the composition AUCs reproduce af3_report_data.json.',
        f'Direction: higher means more interaction, except {", ".join(lower)}, where lower does.',
        f'memorised_control = TRUE: {", ".join(sorted(MEM))}, memorised pipeline controls kept out '
        'of the headline tests.',
        f'Composition baseline, leave-one-out logistic AUC: pooled {loo["pooled_raw"]:.3f}, '
        f'homotypic {loo["homotypic_raw"]:.3f}, heterotypic {loo["heterotypic_raw"]:.3f}. '
        'Scrambles are outside the baseline, so their baseline columns are blank.',
        'Scramble sign test (Metrics): panel metrics count ties as losses; ipsae and lis drop ties '
        '(scramble_ties), as af3_report_stats.py does.',
        "MSA columns count the sequences in the server's a3m files; a homodimer's two chains share "
        'one MSA and one template list.',
        f'PAE and contact-probability matrices (N×N per sample, up to {biggest}×{biggest}) are not '
        'in this workbook. View them with: python3 viewer/serve.py',
    ]
    write_workbook(OUT, sheets, intro, notes)
    print(f'wrote {os.path.basename(OUT)}: '
          + ', '.join(f"{sh['title']} {len(sh['rows'])}" for sh in sheets))
    print(f'all {len(samples)} scored jobs reproduce af3_scored.csv')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'data'))
